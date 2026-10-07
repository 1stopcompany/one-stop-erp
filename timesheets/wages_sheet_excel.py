"""
Excel export of the day-labor wages sheet ("كشف اجور عمال"): the company's columns and formulas on the same formal page as the
employee payroll sheet -- Times New Roman, right-to-left, A4 landscape, the company banner (title in the middle, logo on the left), one
titled table per project (and per sub-group, متفرقة), banded lines, thousands separators, totals, signature boxes and page numbers.

The typed columns (days, paid rate, overtime hours, advances) are values; every other column is a live formula exactly as in the original
sheet (Friday days = days / 6, total days = days + Fridays, daily wage = paid rate x 6 / 7, total = paid rate x days, hour rate =
paid rate / 8 x 1.5, ...), so the file can still be adjusted by hand after export.
"""
import os
from datetime import date

from django.contrib.staticfiles import finders
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.pagebreak import Break

from .payroll_sheet_excel import (
    BAND, COMPANY_AR, COMPANY_EN, FONT, GRID, HAIR, MONEY, NAVY, NET_FILL, TOTAL_FILL, month_label,
)
from django.conf import settings

HEADERS = [
    '#', 'الاســـم', 'طبيعة العمل', 'ايام العمل', 'ايام الجمع', 'مجموع الايام شامل جمع', 'الاجر اليومي', 'الاجر المدفوع',
    'المجموع', 'عدد الساعات الاضافية', 'سعر الساعة', 'قيمة الساعات', 'الاجر المستحق', 'سلف', 'الصافي للدفع', 'رقم الهوية', 'التوقيع',
]
WIDTHS = [4.5, 28, 17, 9, 9, 11.5, 11, 11, 12, 10.5, 10, 11, 12.5, 10, 13, 14, 14]
SUB_FILL = 'DCE6F2'
PAGE_POINTS = 700   # usable height of one printed page in sheet points (the sheet is scaled to fit the width)
LINE_POINTS = 19


def _num(value):
    return float(value)


def _names():
    return getattr(settings, 'WAGES_PREPARED_BY', ''), getattr(settings, 'WAGES_REVIEWED_BY', ''), getattr(settings, 'WAGES_APPROVED_BY', '')


def build_wages_sheet_workbook(sheet):
    wb = Workbook()
    ws = wb.active
    ws.title = f'{sheet["period_start"]:%m.%Y}'
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top, ws.page_margins.bottom = 0.45, 0.6
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = 'صفحة &P من &N'
    ws.oddFooter.center.font = f'{FONT},Regular'
    ws.oddFooter.right.text = 'One Stop ERP'
    ws.oddFooter.left.text = f'{date.today():%d/%m/%Y}'
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[L(i)].width = width
    last_col = len(WIDTHS)

    def font(size=11, bold=False, color='000000'):
        return Font(name=FONT, size=size, bold=bold, color=color)

    centre = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    def merge(r, c1, c2, value, fnt, alignment, height=None):
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(r, c1, value)
        cell.font, cell.alignment = fnt, alignment
        if height:
            ws.row_dimensions[r].height = height
        return cell

    # banner: company, title and the month; the logo on the left
    period_start, period_end = sheet['period_start'], sheet['period_end']
    merge(1, 1, last_col, COMPANY_AR, font(17, True, NAVY), centre, 28)
    merge(2, 1, last_col, COMPANY_EN, font(11, False, '595959'), centre, 17)
    merge(3, 1, last_col, f'كشف اجور عمال - شهر {month_label(period_start)}', font(15, True), centre, 26)
    merge(4, 1, last_col, f'{period_start:%d/%m/%Y}  -  {period_end:%d/%m/%Y}', font(10.5, False, '595959'), centre, 17)
    for c in range(1, last_col + 1):
        ws.cell(4, c).border = Border(bottom=Side(style='medium', color=NAVY))
    logo = finders.find('images/one_stop_logo.png')
    if logo and os.path.exists(logo):
        image = XLImage(logo)
        image.height, image.width = 62, 62 * image.width / image.height if image.height else 62
        ws.add_image(image, 'O1')
    ws.row_dimensions[5].height = 6
    row = 6
    merge(row, 1, last_col, 'إلى : الادارة العامة.   الموضوع: استحقاق صرف الرواتب والأجور.', font(10.5, True), right, 18)
    merge(row + 1, 1, last_col, f'بالإشارة إلى الموضوع أعلاه ، نعلمكم باستحقاق صرف الرواتب والأجور التالية عن شهر {period_start:%m.%Y}',
          font(10.5), right, 18)
    row += 3
    net_cells, used = [], 140

    def money_cols(cell, c):
        if c in (7, 8, 9, 11, 12, 13, 14, 15):
            cell.number_format = MONEY
        elif c in (5, 6):
            cell.number_format = '0.00'   # Friday days and total days (days / 6) -- not 0.666666667

    for section in sheet['sections']:
        name = section['project'].name if section['project'] else '—'
        lines = sum(len(g['rows']) for g in section['groups']) + (2 * len(section['groups']) if section['has_subs'] else 0) + 1
        needed = 22 + 40 + lines * LINE_POINTS + 18
        if used and used + needed > PAGE_POINTS and needed <= PAGE_POINTS:
            ws.row_breaks.append(Break(id=row - 1))
            used = 0
        used = (used + needed) % PAGE_POINTS if needed > PAGE_POINTS else used + needed

        # the project's title bar and the column headers
        merge(row, 1, last_col, f'كشف اجور عمال ({name})', font(12, True, 'FFFFFF'), centre, 22)
        for c in range(1, last_col + 1):
            ws.cell(row, c).fill = PatternFill('solid', fgColor=NAVY)
        for c, text in enumerate(HEADERS, start=1):
            cell = ws.cell(row + 1, c, text)
            cell.font, cell.alignment, cell.border = font(10, True, NAVY), centre, GRID
            cell.fill = PatternFill('solid', fgColor=TOTAL_FILL)
        ws.row_dimensions[row + 1].height = 40
        r = row + 2
        subtotal_rows = []
        first_all = r
        for group in section['groups']:
            if section['has_subs']:
                merge(r, 1, last_col, group['name'] or 'بدون متفرقة', font(11, True, NAVY), right, 19)
                for c in range(1, last_col + 1):
                    ws.cell(r, c).fill = PatternFill('solid', fgColor=SUB_FILL)
                    ws.cell(r, c).border = GRID
                r += 1
            first = r
            for index, line in enumerate(group['rows']):
                values = {
                    1: line['n'], 2: line['worker'].full_name, 3: line['trade'] or '', 4: _num(line['days']),
                    8: _num(line['paid_rate']), 10: _num(line['overtime_hours']) if line['overtime_hours'] else None,
                    14: _num(line['advances']) if line['advances'] else None,
                    16: int(line['national_id']) if str(line['national_id']).isdigit() else line['national_id'],
                }
                formulas = {5: f'=D{r}/6', 6: f'=E{r}+D{r}', 7: f'=H{r}*6/7', 9: f'=H{r}*D{r}', 11: f'=H{r}/8*1.5',
                            12: f'=K{r}*J{r}', 13: f'=L{r}+I{r}', 15: f'=M{r}-N{r}'}
                for c in range(1, last_col + 1):
                    cell = ws.cell(r, c, values.get(c, formulas.get(c)))
                    cell.font = font(11, bold=c in (13, 15))
                    cell.border = GRID
                    cell.alignment = Alignment(horizontal='right' if c in (2, 3) else 'center', vertical='center', wrap_text=c in (2, 3))
                    money_cols(cell, c)
                    if c == 15:
                        cell.fill = PatternFill('solid', fgColor=NET_FILL)
                    elif index % 2:
                        cell.fill = PatternFill('solid', fgColor=BAND)
                r += 1
            last = r - 1
            if section['has_subs']:
                ws.cell(r, 2, f"مجموع {group['name'] or 'بدون متفرقة'}")
                ws.cell(r, 13, f'=SUM(M{first}:M{last})')
                ws.cell(r, 15, f'=SUM(O{first}:O{last})')
                for c in range(1, last_col + 1):
                    cell = ws.cell(r, c)
                    cell.font, cell.fill, cell.border = font(11, True, NAVY), PatternFill('solid', fgColor=SUB_FILL), GRID
                    cell.alignment = Alignment(horizontal='right' if c == 2 else 'center', vertical='center')
                    money_cols(cell, c)
                subtotal_rows.append(r)
                r += 1
        if section['has_subs']:
            ws.cell(r, 13, '=' + '+'.join(f'M{x}' for x in subtotal_rows))
            ws.cell(r, 15, '=' + '+'.join(f'O{x}' for x in subtotal_rows))
            ws.cell(r, 2, 'المجموع النهائي للمتفرقات')
        else:
            ws.cell(r, 13, f'=SUM(M{first_all}:M{r - 1})')
            ws.cell(r, 15, f'=SUM(O{first_all}:O{r - 1})')
            ws.cell(r, 2, 'المجموع')
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.font, cell.fill = font(11.5, True), PatternFill('solid', fgColor=TOTAL_FILL)
            cell.border = Border(top=Side(style='medium', color=NAVY), bottom=Side(style='double', color=NAVY), left=HAIR, right=HAIR)
            cell.alignment = Alignment(horizontal='right' if c == 2 else 'center', vertical='center')
            money_cols(cell, c)
        ws.row_dimensions[r].height = 22
        net_cells.append(f'O{r}')
        row = r + 2

    if sheet['adjustments']:
        merge(row, 1, last_col, 'تسويات على مستوى العامل (بدلات / خصومات / سلف)', font(12, True, 'FFFFFF'), centre, 22)
        for c in range(1, last_col + 1):
            ws.cell(row, c).fill = PatternFill('solid', fgColor=NAVY)
        for c, text in enumerate(['#', 'الاســـم', '', 'بدلات', 'خصومات', 'سلف', 'الأثر على الصافي'], start=1):
            cell = ws.cell(row + 1, c, text)
            cell.font, cell.fill, cell.alignment, cell.border = font(10, True, NAVY), PatternFill('solid', fgColor=TOTAL_FILL), centre, GRID
        r = row + 2
        first = r
        for n, adj in enumerate(sheet['adjustments'], start=1):
            ws.cell(r, 1, n)
            ws.cell(r, 2, adj['worker'].full_name)
            ws.cell(r, 4, _num(adj['allowances']))
            ws.cell(r, 5, _num(adj['deductions']))
            ws.cell(r, 6, _num(adj['advances']))
            ws.cell(r, 7, f'=D{r}-E{r}-F{r}')
            for c in range(1, 8):
                cell = ws.cell(r, c)
                cell.font, cell.border = font(11), GRID
                cell.alignment = Alignment(horizontal='right' if c == 2 else 'center', vertical='center')
                if c >= 4:
                    cell.number_format = MONEY
            r += 1
        ws.cell(r, 2, 'المجموع')
        ws.cell(r, 7, f'=SUM(G{first}:G{r - 1})')
        for c in range(1, 8):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.border = font(11.5, True), PatternFill('solid', fgColor=TOTAL_FILL), GRID
            cell.alignment = Alignment(horizontal='right' if c == 2 else 'center', vertical='center')
        ws.cell(r, 7).number_format = MONEY
        net_cells.append(f'G{r}')
        row = r + 2

    # the grand total
    merge(row, 1, 3, 'المجموع الكلي', font(14, True, NAVY), right, 26)
    total = ws.cell(row, 15, '=' + '+'.join(net_cells) if net_cells else 0)
    total.font, total.fill, total.number_format, total.alignment = font(14, True), PatternFill('solid', fgColor=NET_FILL), MONEY, centre
    total.border = Border(top=Side(style='medium', color=NAVY), bottom=Side(style='double', color=NAVY), left=HAIR, right=HAIR)
    row += 3

    # signatures
    prepared, reviewed, approved = _names()
    block = last_col // 3
    for c1, c2, label, name in [(1, block, 'إعداد', prepared), (block + 1, 2 * block, 'تدقيق', reviewed), (2 * block + 1, last_col, 'اعتماد', approved)]:
        merge(row, c1, c2, label, font(11, True, NAVY), centre, 20)
        merge(row + 1, c1, c2, name or ' ', font(11), centre, 30)
        for c in range(c1, c2 + 1):
            ws.cell(row + 1, c).border = Border(bottom=Side(style='thin', color='7F7F7F'))
    return wb
