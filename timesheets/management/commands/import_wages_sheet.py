"""
Load a month of the company's day-labor wages sheet ("كشف اجور عمال", e.g. FIN-WRK-08-2026.xlsx) into the
Wages Run's Manual Entry data: one section = one project, one line = one worker's days / paid wage / overtime
hours / advance for that month.

    python manage.py import_wages_sheet "FIN-WRK-08-2026.xlsx" --month 2026-08 --dry-run
    python manage.py import_wages_sheet "FIN-WRK-08-2026.xlsx" --month 2026-08

Every sheet in the workbook is read (the main sheet and e.g. "Roof"). A section is attached to an existing
project when its title contains one of the fragments below (or one given with --map "fragment=SYMBOL"), when the
project's name equals the section's name, otherwise a new project is created (status Planning, to be completed
in the Projects screen). Workers are matched by name (same first and last name, spelling variants such as ال ignored); a missing one is added to the day-labor roster with
the sheet's trade, national ID and paid wage. National IDs are unique in the system, so when the sheet gives two
different people the same ID only the first keeps it -- the others are listed so the ID can be fixed by hand.

Safe to re-run: a worker's line for a project/month is updated, not duplicated. A worker whose wages for that
month are already posted is left alone.
"""
import calendar
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from projects.models import Project
from timesheets.models import DailyWorker, DailyWorkerManualEntry, DailyWorkerPayslip

# Section title fragment -> existing project symbol (extend with --map).
DEFAULT_MAP = {'فريديس': 'ORL2', 'طالب عمايرة': 'TAB', 'روف': 'Q1Roof'}


def _squash(text):
    return re.sub(r'\s+', ' ', str(text or '')).strip()


def _section_name(title):
    """'كشف اجور عمال (مشروع عمارة الرجعي  )' -> 'عمارة الرجعي'."""
    inner = re.search(r'\((.*)\)\s*$', title, re.S)
    name = _squash(inner.group(1) if inner else title)
    name = re.sub(r'^(لمشروع|مشروع)\s+', '', name)
    return name


def _tokens(name):
    """Name -> comparable tokens: hamza/ya/ta-marbuta variants folded, and the article "ال" dropped (البستنجي ~ بستنجي)."""
    text = _squash(name).translate(str.maketrans({'أ': 'ا', 'إ': 'ا', 'آ': 'ا', 'ى': 'ي', 'ة': 'ه'}))
    return [t[2:] if t.startswith('ال') and len(t) > 4 else t for t in text.split(' ') if t]


def _same_person(a, b):
    """Same first and last name after folding spelling variants (middle names and 'ال' differences are tolerated)."""
    ta, tb = _tokens(a), _tokens(b)
    return bool(ta and tb and ta[0] == tb[0] and ta[-1] == tb[-1])


def _dec(value):
    try:
        return Decimal(str(value or 0))
    except InvalidOperation:
        return Decimal('0')


def read_sections(path):
    """[(sheet name, section title, [line dicts])] with the cached (calculated) values of the workbook."""
    values = openpyxl.load_workbook(path, data_only=True)
    sections = []
    for ws in values.worksheets:
        current = None
        for row in range(1, ws.max_row + 1):
            a = ws.cell(row, 1).value
            if isinstance(a, str) and 'كشف' in a and 'اجور' in a and '(' in a:
                current = (ws.title, _squash(a), [])
                sections.append(current)
                continue
            name = ws.cell(row, 2).value
            if current is not None and isinstance(a, (int, float)) and name:
                cell = lambda col: ws.cell(row, col).value
                current[2].append({
                    'name': _squash(name), 'trade': _squash(cell(3)), 'days': _dec(cell(4)), 'rate': _dec(cell(8)),
                    'overtime': _dec(cell(10)), 'advances': _dec(cell(14)),
                    'national_id': str(int(cell(16))) if isinstance(cell(16), (int, float)) else _squash(cell(16)),
                })
    return sections


class Command(BaseCommand):
    help = 'Import a day-labor wages sheet (كشف اجور عمال) into the Manual Entry data for one month.'

    def add_arguments(self, parser):
        parser.add_argument('path', help='The .xlsx file')
        parser.add_argument('--month', required=True, help='Month the sheet belongs to, YYYY-MM')
        parser.add_argument('--map', action='append', default=[], metavar='FRAGMENT=SYMBOL',
                            help='Attach sections whose title contains FRAGMENT to the project with this symbol')
        parser.add_argument('--dry-run', action='store_true', help='Show what would happen and change nothing')

    def handle(self, *args, **opts):
        path = Path(opts['path'])
        if not path.exists():
            raise CommandError(f'File not found: {path}')
        try:
            period_start = date.fromisoformat(opts['month'] + '-01')
        except ValueError:
            raise CommandError('--month must look like 2026-08')
        period_end = period_start.replace(day=calendar.monthrange(period_start.year, period_start.month)[1])

        fragments = dict(DEFAULT_MAP)
        for item in opts['map']:
            fragment, _, symbol = item.partition('=')
            fragments[fragment.strip()] = symbol.strip()

        sections = read_sections(path)
        if not sections:
            raise CommandError('No "كشف اجور عمال" sections found in that file.')

        report = {'projects_new': [], 'workers_new': [], 'id_conflicts': [], 'posted': [], 'merged': [], 'lines': 0}
        try:
            with transaction.atomic():
                self._import(sections, fragments, period_start, period_end, path.name, report)
                if opts['dry_run']:
                    raise _DryRun
        except _DryRun:
            self.stdout.write(self.style.WARNING('DRY RUN -- nothing was saved.'))
        self._summary(report, period_start)

    # ------------------------------------------------------------------

    def _project_for(self, title, period_start, report):
        for fragment, symbol in self._fragments.items():
            if fragment in title:
                project = Project.objects.filter(project_symbol=symbol).first()
                if project:
                    return project
        name = _section_name(title)
        project = Project.objects.filter(name=name).first()
        if project:
            return project
        n = 1
        while Project.objects.filter(project_symbol=f'WRK-{n:02d}').exists() or Project.objects.filter(contract_number=f'WRK-{n:02d}').exists():
            n += 1
        project = Project.objects.create(
            name=name, project_symbol=f'WRK-{n:02d}', contract_number=f'WRK-{n:02d}',
            client_name='(to be completed)', start_date=period_start, status='planning',
        )
        report['projects_new'].append(project)
        return project

    def _worker_for(self, line, report):
        worker = DailyWorker.objects.filter(full_name=line['name']).first()
        if worker:
            return worker
        similar = [w for w in DailyWorker.objects.all() if _same_person(w.full_name, line['name'])]
        if len(similar) > 1:
            # Several share first/last name: prefer the one with the same national ID, else the one sharing most name parts.
            by_id = [w for w in similar if line['national_id'] and w.national_id == line['national_id']]
            if by_id:
                similar = by_id
            else:
                overlap = {w.pk: len(set(_tokens(w.full_name)) & set(_tokens(line['name']))) for w in similar}
                best = max(overlap.values())
                similar = [w for w in similar if overlap[w.pk] == best]
        if len(similar) == 1:
            return similar[0]
        national_id = line['national_id'] or None
        if national_id and DailyWorker.objects.filter(national_id=national_id).exists():
            holder = DailyWorker.objects.get(national_id=national_id)
            report['id_conflicts'].append((line['name'], national_id, holder.full_name))
            national_id = None
        worker = DailyWorker.objects.create(
            full_name=line['name'], trade=line['trade'], national_id=national_id, daily_rate=line['rate'] or Decimal('0'),
        )
        report['workers_new'].append(worker)
        return worker

    def _import(self, sections, fragments, period_start, period_end, filename, report):
        self._fragments = fragments
        posted = set(
            DailyWorkerPayslip.objects.filter(period_start=period_start, period_end=period_end, status='posted')
            .values_list('worker_id', flat=True)
        )
        for sheet, title, lines in sections:
            project = self._project_for(title, period_start, report)
            merged = defaultdict(lambda: {'days': Decimal('0'), 'overtime': Decimal('0'), 'advances': Decimal('0'), 'rate': None})
            workers = {}
            for line in lines:
                worker = self._worker_for(line, report)
                workers[worker.pk] = worker
                bucket = merged[worker.pk]
                if bucket['rate'] is not None:
                    report['merged'].append((worker.full_name, project.name))
                bucket['days'] += line['days']
                bucket['overtime'] += line['overtime']
                bucket['advances'] += line['advances']
                bucket['rate'] = bucket['rate'] or line['rate'] or worker.daily_rate
            for worker_id, data in merged.items():
                worker = workers[worker_id]
                if worker_id in posted:
                    report['posted'].append((worker.full_name, project.name))
                    continue
                if data['days'] == 0 and data['overtime'] == 0 and data['advances'] == 0:
                    continue
                DailyWorkerManualEntry.objects.update_or_create(
                    worker=worker, project=project, period_start=period_start, sub_name="",
                    defaults={
                        'days': data['days'], 'daily_rate': data['rate'], 'overtime_hours': data['overtime'],
                        'advances': data['advances'], 'note': f'Imported from {filename} [{sheet}]',
                    },
                )
                report['lines'] += 1

    def _summary(self, report, period_start):
        w = self.stdout.write
        w(f'Lines imported for {period_start:%Y-%m}: {report["lines"]}')
        if report['projects_new']:
            w('New projects (status Planning -- complete their details in Projects): ' + '; '.join(
                f'{p.project_symbol} {p.name}' for p in report['projects_new']))
        if report['workers_new']:
            w(f'New workers added to the roster: {len(report["workers_new"])}')
        for name, nid, holder in report['id_conflicts']:
            w(self.style.WARNING(f'  National ID {nid} for "{name}" is already used by "{holder}" -- left blank for {name}; fix by hand.'))
        for name, project in report['merged']:
            w(self.style.WARNING(f'  {name} appears twice under "{project}" -- days/overtime/advances were added together.'))
        for name, project in report['posted']:
            w(self.style.WARNING(f'  {name}: this month is already posted -- skipped ({project}).'))


class _DryRun(Exception):
    pass
