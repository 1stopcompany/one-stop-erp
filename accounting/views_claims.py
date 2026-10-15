from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Q
from django.http import Http404, HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from procurement.models import Vendor
from projects.models import Project

from . import payment_claims as service
from .claim_pdf import generate_payment_claim_pdf
from .models import REQUIREMENTS, PaymentClaim


def _text(error):
    return '; '.join(error.messages) if isinstance(error, ValidationError) else str(error)


def _lines(post):
    return list(zip(post.getlist('line_statement'), post.getlist('line_details'), post.getlist('line_amount')))


def _get_claim(request, pk):
    return get_object_or_404(service.visible_claims(request.user), pk=pk)


# ------------------------------------------------------------------ list
@login_required
def claim_list(request):
    qs = service.visible_claims(request.user)
    q = request.GET.get('q', '').strip()
    status = request.GET.get('status', '')
    project = request.GET.get('project', '')
    if q:
        qs = qs.filter(Q(number__icontains=q) | Q(beneficiary_name__icontains=q) | Q(nature_of_work__icontains=q))
    if status == 'open':
        qs = qs.exclude(status__in=(PaymentClaim.PAID,))
    elif status in dict(PaymentClaim.STATUS_CHOICES):
        qs = qs.filter(status=status)
    if project == 'main':
        qs = qs.filter(project__isnull=True)
    elif project.isdigit():
        qs = qs.filter(project_id=project)
    claims = list(qs.prefetch_related('lines')[:400])
    waiting = [c for c in claims if service.can_sign(request.user, c)]
    return render(request, 'accounting/claims/list.html', {
        'claims': claims, 'q': q, 'status': status, 'project': project, 'statuses': PaymentClaim.STATUS_CHOICES,
        'projects': Project.objects.filter(payment_claims__isnull=False).distinct().order_by('project_symbol'),
        'can_create': service.can_create_any(request.user), 'waiting_for_me': waiting,
        'to_pay': [c for c in claims if service.can_pay(request.user, c)],
        'sum_net': sum((c.net_due for c in claims), Decimal('0')),
    })


# ------------------------------------------------------------------ create / edit
FORM_FIELDS = ('vendor', 'beneficiary_name', 'claim_type', 'claim_date', 'nature_of_work', 'original_contract_value', 'current_contract_value',
               'previous_payments', 'notes', 'deduction_agreed', 'deduction_source_retention', 'deduction_other', 'project')


def _form_context(request, claim, post=None):
    """What the form shows: what was just posted (after a refusal), else the claim's own values, else blanks."""
    if post is not None and hasattr(post, 'getlist'):
        values = {name: post.get(name, '') for name in FORM_FIELDS}
        ticked = {key for key, _label in REQUIREMENTS if post.get(f'req_{key}')}
        rows = [{'statement': a, 'details': b, 'amount': c} for a, b, c in _lines(post)]
    elif claim is not None:
        values = {name: ('' if getattr(claim, name, None) is None else getattr(claim, name)) for name in FORM_FIELDS if name not in ('vendor', 'project')}
        values['vendor'] = claim.vendor_id or ''
        values['project'] = claim.project_id or 'main'
        ticked = set(claim.requirements or [])
        rows = [{'statement': l.statement, 'details': l.details, 'amount': l.amount} for l in claim.lines.all()]
    else:
        values = {name: '' for name in FORM_FIELDS}
        values.update(post or {})
        values['claim_type'] = 'interim'
        values['claim_date'] = timezone.localdate().isoformat()
        for name in ('deduction_agreed', 'deduction_source_retention', 'deduction_other'):
            values[name] = '0'
        ticked, rows = set(), []
    values['claim_date'] = str(values['claim_date'])[:10]
    return {
        'claim': claim, 'v': values, 'ticked': ticked, 'line_rows': rows or [{'statement': 'Claim details / تفاصيل المطالبة', 'details': '', 'amount': ''}],
        'vendors': Vendor.objects.filter(is_active=True).order_by('name'), 'requirements': REQUIREMENTS,
        'projects': service.projects_for_new_claim(request.user), 'can_main_office': service.can_create(request.user, None),
        'types': PaymentClaim.TYPE_CHOICES,
    }


@login_required
def claim_new(request):
    if not service.can_create_any(request.user):
        return HttpResponseForbidden('Only the project manager prepares a payment claim.')
    if request.method == 'POST':
        project = Project.objects.filter(pk=request.POST.get('project')).first() if request.POST.get('project') not in ('', 'main', None) else None
        try:
            claim = service.create_claim(request.user, project, request.POST, _lines(request.POST))
        except (ValidationError, PermissionDenied) as error:
            messages.error(request, _text(error))
            return render(request, 'accounting/claims/form.html', _form_context(request, None, request.POST))
        messages.success(request, f'Claim {claim.number} was created as a draft. Check it, then press "Sign and send".')
        return redirect('accounting:claim_detail', pk=claim.pk)
    initial = None
    if request.GET.get('project', '').isdigit():
        initial = {'project': request.GET['project']}
    return render(request, 'accounting/claims/form.html', _form_context(request, None, initial))


@login_required
def claim_edit(request, pk):
    claim = _get_claim(request, pk)
    if not service.can_edit(request.user, claim):
        messages.error(request, 'This claim cannot be edited.')
        return redirect('accounting:claim_detail', pk=pk)
    if request.method == 'POST':
        try:
            service.update_claim(request.user, claim, request.POST, _lines(request.POST))
            messages.success(request, 'Saved.')
            return redirect('accounting:claim_detail', pk=pk)
        except (ValidationError, PermissionDenied) as error:
            messages.error(request, _text(error))
    return render(request, 'accounting/claims/form.html', _form_context(request, claim))


# ------------------------------------------------------------------ one claim
@login_required
def claim_detail(request, pk):
    claim = _get_claim(request, pk)
    user = request.user
    signatures = [('Project manager (prepared)', claim.submitted_by, claim.submitted_at), ('Finance', claim.finance_by, claim.finance_at),
                  ('Executive director', claim.executive_by, claim.executive_at), ('General manager', claim.general_by, claim.general_at)]
    return render(request, 'accounting/claims/detail.html', {
        'claim': claim, 'lines': claim.lines.all(), 'signatures': signatures, 'step_label': service.STEP_LABELS.get(claim.next_signer, ''),
        'can_edit': service.can_edit(user, claim), 'can_submit': service.can_submit(user, claim), 'can_sign': service.can_sign(user, claim),
        'can_pay': service.can_pay(user, claim), 'can_delete': service.can_delete(user, claim),
    })


@login_required
@require_POST
def claim_action(request, pk, action):
    claim = _get_claim(request, pk)
    post = request.POST
    try:
        if action == 'submit':
            service.submit(request.user, claim)
            messages.success(request, 'Signed and sent: it now waits for the finance.')
        elif action == 'sign':
            done = service.sign(request.user, claim)
            messages.success(request, 'Signed.' + (' The claim is approved and ready to pay.' if done.status == PaymentClaim.APPROVED else ''))
        elif action == 'reject':
            service.reject(request.user, claim, post.get('reason', ''))
            messages.warning(request, 'The claim was sent back to its preparer.')
        elif action == 'pay':
            service.pay(request.user, claim, post)
            messages.success(request, 'The payment was recorded.')
        elif action == 'reopen':
            service.reopen(request.user, claim)
            messages.warning(request, 'The payment was reopened.')
        elif action == 'delete':
            service.delete_claim(request.user, claim)
            messages.info(request, 'The draft was deleted.')
            return redirect('accounting:claim_list')
        else:
            raise Http404
    except (ValidationError, PermissionDenied) as error:
        messages.error(request, _text(error))
    return redirect('accounting:claim_detail', pk=pk)


@login_required
def claim_pdf(request, pk):
    claim = _get_claim(request, pk)
    response = HttpResponse(generate_payment_claim_pdf(claim), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{claim.number}.pdf"'
    return response


@login_required
@require_POST
def claim_from_equipment(request, pk):
    from equipment.models import EquipmentClaim
    equipment = get_object_or_404(EquipmentClaim, pk=pk)
    try:
        claim = service.from_equipment_claim(request.user, equipment)
        messages.success(request, f'Claim {claim.number} was created from the rental claim. Check it, then sign and send it.')
        return redirect('accounting:claim_detail', pk=claim.pk)
    except (ValidationError, PermissionDenied) as error:
        messages.error(request, _text(error))
        return redirect(request.META.get('HTTP_REFERER') or 'accounting:payables')


# ------------------------------------------------------------------ the claims report (the "payment sheet")
@login_required
def claims_report(request):
    qs = service.visible_claims(request.user).prefetch_related('lines')
    status = request.GET.get('status', '')
    project = request.GET.get('project', '')
    date_from, date_to = request.GET.get('from', ''), request.GET.get('to', '')
    vendor = request.GET.get('vendor', '').strip()
    if status == 'unpaid':
        qs = qs.exclude(status=PaymentClaim.PAID)
    elif status in dict(PaymentClaim.STATUS_CHOICES):
        qs = qs.filter(status=status)
    if project == 'main':
        qs = qs.filter(project__isnull=True)
    elif project.isdigit():
        qs = qs.filter(project_id=project)
    if date_from:
        qs = qs.filter(claim_date__gte=date_from)
    if date_to:
        qs = qs.filter(claim_date__lte=date_to)
    if vendor:
        qs = qs.filter(beneficiary_name__icontains=vendor)
    claims = list(qs.order_by('claim_date', 'id')[:2000])
    rows = [{'claim': c, 'net': c.net_due, 'payment': _payment_text(c)} for c in claims]
    total = sum((r['net'] for r in rows), Decimal('0'))
    paid_total = sum((r['net'] for r in rows if r['claim'].status == PaymentClaim.PAID), Decimal('0'))
    if request.GET.get('format') == 'xlsx':
        return _report_xlsx(rows, total, paid_total)
    return render(request, 'accounting/claims/report.html', {
        'rows': rows, 'total': total, 'paid_total': paid_total, 'unpaid_total': total - paid_total, 'statuses': PaymentClaim.STATUS_CHOICES,
        'projects': Project.objects.filter(payment_claims__isnull=False).distinct().order_by('project_symbol'),
        'status': status, 'project': project, 'date_from': date_from, 'date_to': date_to, 'vendor': vendor,
    })


def _payment_text(claim):
    if claim.status != PaymentClaim.PAID:
        return claim.get_status_display()
    parts = []
    if claim.cash_amount:
        parts.append(f'cash {claim.cash_amount:,.2f}' + (f' ({claim.cash_date:%d/%m/%Y})' if claim.cash_date else ''))
    if claim.cheque_amount:
        parts.append(f'cheque {claim.cheque_amount:,.2f}' + (f' no. {claim.cheque_no}' if claim.cheque_no else '') + (f' due {claim.cheque_due_date:%d/%m/%Y}' if claim.cheque_due_date else ''))
    if claim.payment_note:
        parts.append(claim.payment_note)
    return 'Paid: ' + '; '.join(parts) if parts else 'Paid'


def _report_xlsx(rows, total, paid_total):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Payment claims'
    ws.sheet_view.rightToLeft = True
    ws.append(['تقرير المطالبات المالية'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    headers = ['المورد / المستفيد', 'المشروع', 'تاريخ تقديم المطالبة', 'رقم المطالبة', 'المبلغ المستحق', 'الحالة / الدفع']
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DBE8F5')
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        claim = row['claim']
        ws.append([claim.beneficiary_name, claim.project.project_symbol if claim.project_id else 'Main office', claim.claim_date, claim.number, float(row['net']), _payment_text(claim)])
        ws.cell(row=ws.max_row, column=3).number_format = 'DD/MM/YYYY'
        ws.cell(row=ws.max_row, column=5).number_format = '#,##0.00'
    ws.append([])
    ws.append(['', '', '', 'الإجمالي', float(total)])
    ws.append(['', '', '', 'المدفوع', float(paid_total)])
    ws.append(['', '', '', 'غير المدفوع', float(total - paid_total)])
    for column, width in zip('ABCDEF', (34, 14, 18, 24, 16, 52)):
        ws.column_dimensions[column].width = width
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="payment-claims-{timezone.localdate():%Y-%m-%d}.xlsx"'
    return response
