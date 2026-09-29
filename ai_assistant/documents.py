"""Turn an uploaded file into the content blocks Claude can read."""
import base64
import csv
import io
import os

from django.conf import settings

PDF = {".pdf"}
IMAGES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}
SHEETS = {".xlsx", ".xlsm"}
TEXT = {".csv", ".txt"}
CAD = {".dwg", ".dxf", ".rvt", ".ifc", ".zip", ".rar"}

MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 600_000


class UnsupportedSource(ValueError):
    """The file can't be read by the assistant; the message tells the user what to do instead."""


def source_of(run):
    """(file name, open file object) for whatever the run was pointed at."""
    for field in (run.source_file, getattr(run.source_tender, "document", None), getattr(run.source_revision, "file", None)):
        if field:
            return os.path.basename(field.name), field
    raise UnsupportedSource("This run has no file attached.")


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("ascii")


def _sheet_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            if any(cell not in (None, "") for cell in row):
                parts.append(" | ".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(parts)


def _load(name: str, file_obj):
    """(extension, bytes) after the checks every provider shares."""
    ext = os.path.splitext(name)[1].lower()
    if ext in CAD:
        raise UnsupportedSource(
            f"{ext.upper()} files can't be read directly. Export the drawing to PDF from the CAD program and upload the PDF."
        )
    max_bytes = settings.AI_MAX_UPLOAD_MB * 1024 * 1024
    file_obj.open("rb")
    try:
        data = file_obj.read()
    finally:
        file_obj.close()
    if len(data) > max_bytes:
        raise UnsupportedSource(f"The file is {len(data) / 1024 / 1024:.0f} MB; the limit is {settings.AI_MAX_UPLOAD_MB} MB. Split it into parts.")
    return ext, data


def content_blocks(name: str, file_obj) -> list:
    """The document as Claude API content blocks. Raises UnsupportedSource with a plain-language reason."""
    ext, data = _load(name, file_obj)

    if ext in PDF:
        _check_pdf_pages(data)
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(data)}, "title": name}]
    if ext in IMAGES:
        if len(data) > MAX_IMAGE_BYTES:
            raise UnsupportedSource("The image is larger than 5 MB. Save it as a PDF or a smaller image.")
        return [{"type": "image", "source": {"type": "base64", "media_type": IMAGES[ext], "data": _b64(data)}}]
    if ext in SHEETS:
        text = _sheet_text(data)
    elif ext in TEXT:
        text = data.decode("utf-8-sig", errors="replace")
        if ext == ".csv":
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    else:
        raise UnsupportedSource(f"'{ext or 'this'}' files aren't supported. Use PDF, an image, Excel (.xlsx), CSV or text.")
    if len(text) > MAX_TEXT_CHARS:
        raise UnsupportedSource("The sheet is too large for one pass. Split it into parts.")
    return [{"type": "text", "text": f"Contents of {name}:\n\n{text}"}]


def _check_pdf_pages(data: bytes) -> None:
    try:
        import pymupdf as fitz  # PyMuPDF, only used to count pages
    except ImportError:
        try:
            import fitz
        except ImportError:
            return
    try:
        pages = len(fitz.open(stream=data, filetype="pdf"))
    except Exception:
        raise UnsupportedSource("This PDF is damaged or password-protected and can't be opened.")
    if pages > settings.AI_MAX_PDF_PAGES:
        raise UnsupportedSource(f"The PDF has {pages} pages; the limit is {settings.AI_MAX_PDF_PAGES}. Split it into parts.")


# ------------------------------------------------------------------ local (Ollama) reading

SCANNED_PAGE_CHARS = 30      # a page with fewer characters of text than this is treated as a picture (scanned)
RENDER_DPI = 150


def _pymupdf():
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf
        except ImportError:
            raise UnsupportedSource("Reading PDFs with the local AI needs the PyMuPDF package: pip install pymupdf")
    return pymupdf


def local_pages(name: str, file_obj) -> list:
    """
    The document as pages for a model that can't take PDFs: [{"page": n, "text": str, "image": base64 PNG or None}].
    A page carries text, or (scanned pages, photos, or when OLLAMA_PDF_MODE=images) a picture for a vision model.
    Raises UnsupportedSource with a plain-language reason.
    """
    ext, data = _load(name, file_obj)
    vision = bool(settings.OLLAMA_VISION_MODEL)
    mode = settings.OLLAMA_PDF_MODE

    if ext in IMAGES:
        if len(data) > MAX_IMAGE_BYTES:
            raise UnsupportedSource("The image is larger than 5 MB. Save it as a PDF or a smaller image.")
        if not vision:
            raise UnsupportedSource("Reading images needs a vision model: set OLLAMA_VISION_MODEL in the .env file (and pull it with Ollama).")
        return [{"page": 1, "text": "", "image": _b64(data)}]

    if ext in PDF:
        if mode == "images" and not vision:
            raise UnsupportedSource("OLLAMA_PDF_MODE=images needs OLLAMA_VISION_MODEL to be set.")
        pymupdf = _pymupdf()
        try:
            pdf = pymupdf.open(stream=data, filetype="pdf")
        except Exception:
            raise UnsupportedSource("This PDF is damaged or password-protected and can't be opened.")
        if len(pdf) > settings.AI_MAX_PDF_PAGES:
            raise UnsupportedSource(f"The PDF has {len(pdf)} pages; the limit is {settings.AI_MAX_PDF_PAGES}. Split it into parts.")
        pages, pictures = [], 0
        for number, page in enumerate(pdf, start=1):
            text = page.get_text().strip()
            scanned = len(text) < SCANNED_PAGE_CHARS
            if vision and (mode == "images" or (mode == "auto" and scanned)):
                pictures += 1
                if pictures > settings.OLLAMA_MAX_IMAGE_PAGES:
                    raise UnsupportedSource(
                        f"More than {settings.OLLAMA_MAX_IMAGE_PAGES} pages need to be read as pictures. Split the PDF into parts "
                        "or raise OLLAMA_MAX_IMAGE_PAGES."
                    )
                png = page.get_pixmap(dpi=RENDER_DPI).tobytes("png")
                pages.append({"page": number, "text": "", "image": _b64(png)})
            elif len(text) < SCANNED_PAGE_CHARS:
                pages.append({"page": number, "text": "", "image": None})     # blank or scanned, and no vision model
            else:
                pages.append({"page": number, "text": text, "image": None})
        if not any(p["text"] or p["image"] for p in pages):
            raise UnsupportedSource(
                "This PDF has no readable text (it looks scanned). Set OLLAMA_VISION_MODEL in the .env file so pages can be read as pictures."
            )
        return pages

    if ext in SHEETS:
        text = _sheet_text(data)
    elif ext in TEXT:
        text = data.decode("utf-8-sig", errors="replace")
        if ext == ".csv":
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    else:
        raise UnsupportedSource(f"'{ext or 'this'}' files aren't supported. Use PDF, an image, Excel (.xlsx), CSV or text.")
    return [{"page": 1, "text": text, "image": None}]
