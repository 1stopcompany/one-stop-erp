"""
Views and API endpoints for Material Tracking System
"""

from django.shortcuts import render, redirect, get_object_or_404
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.urls import reverse_lazy
from django.db.models import Sum, Count, Q, F
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from datetime import datetime, timedelta
from decimal import Decimal

from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models_materials import (
    MaterialCategory, MaterialType, Project, MaterialEntry,
    MonthlyMaterialReport, MonthlyMaterialReportItem,
    MaterialBudget, MaterialBudgetItem, MaterialSupplier,
    MaterialInvoice, MaterialInvoiceItem
)
from .serializers_materials import (
    MaterialCategorySerializer, MaterialTypeSerializer, ProjectSerializer,
    MaterialEntrySerializer, MonthlyMaterialReportSerializer,
    MaterialBudgetSerializer, MaterialSupplierSerializer,
    MaterialInvoiceSerializer, MaterialReportSummarySerializer
)
from .forms_materials import (
    MaterialCategoryForm, MaterialTypeForm, ProjectForm, MaterialEntryForm,
    MonthlyMaterialReportForm, MaterialBudgetForm, MaterialSupplierForm,
    MaterialInvoiceForm, MaterialReportFilterForm
)


# ============================================================================
# WEB VIEWS
# ============================================================================

class ProjectListView(LoginRequiredMixin, ListView):
    """List all projects"""
    model = Project
    template_name = 'procurement/project_list.html'
    context_object_name = 'projects'
    paginate_by = 20
    
    def get_queryset(self):
        return Project.objects.all().order_by('-created_at')


class ProjectDetailView(LoginRequiredMixin, DetailView):
    """Project detail with material summary"""
    model = Project
    template_name = 'procurement/project_detail.html'
    context_object_name = 'project'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        project = self.get_object()
        
        # Material statistics
        entries = MaterialEntry.objects.filter(project=project)
        context['total_entries'] = entries.count()
        context['total_quantity'] = entries.aggregate(Sum('quantity'))['quantity__sum'] or 0
        context['total_cost'] = entries.aggregate(Sum('total_cost'))['total_cost__sum'] or 0
        
        # Monthly reports
        context['monthly_reports'] = MonthlyMaterialReport.objects.filter(
            project=project
        ).order_by('-report_month')[:6]
        
        # Budget info
        try:
            context['budget'] = MaterialBudget.objects.get(project=project)
        except MaterialBudget.DoesNotExist:
            context['budget'] = None
        
        return context


class ProjectCreateView(LoginRequiredMixin, CreateView):
    """Create new project"""
    model = Project
    form_class = ProjectForm
    template_name = 'procurement/project_form.html'
    success_url = reverse_lazy('procurement:project_list')


class MaterialEntryListView(LoginRequiredMixin, ListView):
    """List material entries for a project"""
    model = MaterialEntry
    template_name = 'procurement/material_entry_list.html'
    context_object_name = 'entries'
    paginate_by = 50
    
    def get_queryset(self):
        project_id = self.kwargs.get('project_id')
        return MaterialEntry.objects.filter(
            project_id=project_id
        ).select_related('material_type', 'project').order_by('-entry_date')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        project_id = self.kwargs.get('project_id')
        context['project'] = Project.objects.get(id=project_id)
        context['filter_form'] = MaterialReportFilterForm()
        return context


class MaterialEntryCreateView(LoginRequiredMixin, CreateView):
    """Create new material entry"""
    model = MaterialEntry
    form_class = MaterialEntryForm
    template_name = 'procurement/material_entry_form.html'
    
    def get_success_url(self):
        return reverse_lazy('procurement:material_entry_list', kwargs={'project_id': self.object.project.id})
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['project'] = Project.objects.get(id=self.kwargs.get('project_id'))
        return context


class MaterialEntryUpdateView(LoginRequiredMixin, UpdateView):
    """Update material entry"""
    model = MaterialEntry
    form_class = MaterialEntryForm
    template_name = 'procurement/material_entry_form.html'
    
    def get_success_url(self):
        return reverse_lazy('procurement:material_entry_list', kwargs={'project_id': self.object.project.id})


class MonthlyReportListView(LoginRequiredMixin, ListView):
    """List monthly material reports"""
    model = MonthlyMaterialReport
    template_name = 'procurement/monthly_report_list.html'
    context_object_name = 'reports'
    paginate_by = 20
    
    def get_queryset(self):
        project_id = self.kwargs.get('project_id')
        return MonthlyMaterialReport.objects.filter(
            project_id=project_id
        ).order_by('-report_month')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['project'] = Project.objects.get(id=self.kwargs.get('project_id'))
        return context


class MonthlyReportDetailView(LoginRequiredMixin, DetailView):
    """Monthly report detail with items"""
    model = MonthlyMaterialReport
    template_name = 'procurement/monthly_report_detail.html'
    context_object_name = 'report'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        report = self.get_object()
        context['items'] = report.items.all().order_by('material_type__category', 'material_type__name')
        return context


class MonthlyReportCreateView(LoginRequiredMixin, CreateView):
    """Create monthly report"""
    model = MonthlyMaterialReport
    form_class = MonthlyMaterialReportForm
    template_name = 'procurement/monthly_report_form.html'
    
    def get_success_url(self):
        return reverse_lazy('procurement:monthly_report_detail', kwargs={'pk': self.object.id})


class MaterialBudgetDetailView(LoginRequiredMixin, DetailView):
    """Material budget detail"""
    model = MaterialBudget
    template_name = 'procurement/material_budget_detail.html'
    context_object_name = 'budget'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        budget = self.get_object()
        context['items'] = budget.items.all()
        context['spent_amount'] = budget.get_spent_amount()
        context['remaining_budget'] = budget.get_remaining_budget()
        context['budget_percentage'] = budget.get_budget_percentage()
        return context


# ============================================================================
# REST API VIEWSETS
# ============================================================================

class MaterialCategoryViewSet(viewsets.ModelViewSet):
    """API endpoint for Material Categories"""
    queryset = MaterialCategory.objects.all()
    serializer_class = MaterialCategorySerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name']
    ordering_fields = ['name', 'created_at']
    ordering = ['name']


class MaterialTypeViewSet(viewsets.ModelViewSet):
    """API endpoint for Material Types"""
    queryset = MaterialType.objects.all()
    serializer_class = MaterialTypeSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name', 'category__name']
    ordering_fields = ['name', 'category', 'unit_cost']
    ordering = ['category', 'name']


class ProjectViewSet(viewsets.ModelViewSet):
    """API endpoint for Projects"""
    queryset = Project.objects.all()
    serializer_class = ProjectSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['code', 'name', 'location']
    ordering_fields = ['code', 'name', 'start_date']
    ordering = ['-created_at']
    
    @action(detail=True, methods=['get'])
    def summary(self, request, pk=None):
        """Get project material summary"""
        project = self.get_object()
        entries = MaterialEntry.objects.filter(project=project)
        
        summary = {
            'project_code': project.code,
            'project_name': project.name,
            'total_entries': entries.count(),
            'total_quantity': float(entries.aggregate(Sum('quantity'))['quantity__sum'] or 0),
            'total_cost': float(entries.aggregate(Sum('total_cost'))['total_cost__sum'] or 0),
            'by_category': []
        }
        
        # Group by category
        categories = MaterialCategory.objects.filter(material_types__entries__project=project).distinct()
        for category in categories:
            cat_entries = entries.filter(material_type__category=category)
            summary['by_category'].append({
                'category': category.name,
                'unit': category.unit,
                'quantity': float(cat_entries.aggregate(Sum('quantity'))['quantity__sum'] or 0),
                'cost': float(cat_entries.aggregate(Sum('total_cost'))['total_cost__sum'] or 0),
            })
        
        return Response(summary)


class MaterialEntryViewSet(viewsets.ModelViewSet):
    """API endpoint for Material Entries"""
    queryset = MaterialEntry.objects.all()
    serializer_class = MaterialEntrySerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['material_type__name', 'project__code', 'invoice_number']
    ordering_fields = ['entry_date', 'total_cost']
    ordering = ['-entry_date']
    
    def get_queryset(self):
        queryset = MaterialEntry.objects.all()
        project_id = self.request.query_params.get('project_id')
        if project_id:
            queryset = queryset.filter(project_id=project_id)
        return queryset


class MonthlyMaterialReportViewSet(viewsets.ModelViewSet):
    """API endpoint for Monthly Material Reports"""
    queryset = MonthlyMaterialReport.objects.all()
    serializer_class = MonthlyMaterialReportSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['report_month']
    ordering = ['-report_month']
    
    @action(detail=True, methods=['post'])
    def generate(self, request, pk=None):
        """Generate monthly report from entries"""
        report = self.get_object()
        entries = report.generate_report()
        
        # Create report items
        report.items.all().delete()
        
        for category in MaterialCategory.objects.all():
            cat_entries = entries.filter(material_type__category=category)
            if cat_entries.exists():
                for material_type in MaterialType.objects.filter(category=category):
                    mat_entries = cat_entries.filter(material_type=material_type)
                    if mat_entries.exists():
                        quantity = mat_entries.aggregate(Sum('quantity'))['quantity__sum'] or 0
                        total_cost = mat_entries.aggregate(Sum('total_cost'))['total_cost__sum'] or 0
                        
                        MonthlyMaterialReportItem.objects.create(
                            report=report,
                            material_type=material_type,
                            quantity=quantity,
                            unit_cost=material_type.unit_cost,
                            total_cost=total_cost,
                            entry_count=mat_entries.count()
                        )
        
        return Response({'status': 'Report generated successfully'})
    
    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        """Approve monthly report"""
        report = self.get_object()
        report.is_approved = True
        report.approved_by = request.user.username
        report.approved_date = timezone.now()
        report.save()
        return Response({'status': 'Report approved'})


class MaterialBudgetViewSet(viewsets.ModelViewSet):
    """API endpoint for Material Budgets"""
    queryset = MaterialBudget.objects.all()
    serializer_class = MaterialBudgetSerializer
    permission_classes = [IsAuthenticated]
    
    @action(detail=True, methods=['get'])
    def status(self, request, pk=None):
        """Get budget status"""
        budget = self.get_object()
        return Response({
            'project_code': budget.project.code,
            'total_budget': float(budget.total_budget),
            'spent_amount': float(budget.get_spent_amount()),
            'remaining_budget': float(budget.get_remaining_budget()),
            'budget_percentage': round(budget.get_budget_percentage(), 2),
        })


class MaterialSupplierViewSet(viewsets.ModelViewSet):
    """API endpoint for Material Suppliers"""
    queryset = MaterialSupplier.objects.all()
    serializer_class = MaterialSupplierSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name', 'city', 'country']
    ordering_fields = ['name']
    ordering = ['name']


class MaterialInvoiceViewSet(viewsets.ModelViewSet):
    """API endpoint for Material Invoices"""
    queryset = MaterialInvoice.objects.all()
    serializer_class = MaterialInvoiceSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['invoice_number', 'supplier__name']
    ordering_fields = ['invoice_date', 'total_amount']
    ordering = ['-invoice_date']


# ============================================================================
# EXPORT FUNCTIONS
# ============================================================================

def export_monthly_report_excel(request, report_id):
    """Export monthly report to Excel"""
    report = get_object_or_404(MonthlyMaterialReport, id=report_id)
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Report {report.report_month.strftime('%Y-%m')}"
    
    # Set column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    
    # Header
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Material Type', 'Unit', 'Quantity', 'Unit Cost', 'Total Cost', 'Entries']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Data
    row = 2
    total_quantity = 0
    total_cost = 0
    
    for item in report.items.all():
        ws.cell(row=row, column=1).value = item.material_type.name
        ws.cell(row=row, column=2).value = item.material_type.unit
        ws.cell(row=row, column=3).value = float(item.quantity)
        ws.cell(row=row, column=4).value = float(item.unit_cost)
        ws.cell(row=row, column=5).value = float(item.total_cost)
        ws.cell(row=row, column=6).value = item.entry_count
        
        total_quantity += float(item.quantity)
        total_cost += float(item.total_cost)
        row += 1
    
    # Totals
    total_row = row + 1
    ws.cell(row=total_row, column=1).value = "Total"
    ws.cell(row=total_row, column=3).value = total_quantity
    ws.cell(row=total_row, column=5).value = total_cost